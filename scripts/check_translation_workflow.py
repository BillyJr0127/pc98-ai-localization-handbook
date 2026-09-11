"""Executable synthetic XLSX regression; no commercial data or game build is used."""
import copy
import json
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from engine_adapter import SyntheticAdapter, digest
from translation_exchange import (SHEETS, export_workbooks, import_candidate,
                                  load_state, read_json, write_new_json)

ROOT = Path(__file__).resolve().parents[1]


def reject(function, contains=None):
    try:
        function()
    except (ValueError, OSError) as error:
        if contains:
            assert contains in str(error), str(error)
        return
    raise AssertionError('Expected rejection')


def main():
    adapter = SyntheticAdapter(read_json(ROOT/'examples/translation-catalog.json'))
    original = copy.deepcopy(adapter.catalog)
    first, second, third, fourth = [item['id'] for item in adapter.catalog['occurrences']]
    state = load_state(adapter)
    state['languages']['zh-TW'][third] = '燈亮著。\n桌面是空的。'
    state['languages']['yue-HK'][first] = '你好，{player}。'
    with tempfile.TemporaryDirectory(prefix='pc98-excel-example-') as folder:
        temp = Path(folder)
        # CLI smoke: independent exporter and importer operate on real files.
        def cli(script, *args, ok=True):
            result = subprocess.run([sys.executable, '-X', 'utf8', str(ROOT/'scripts'/script), *map(str, args)],
                                    capture_output=True, text=True, encoding='utf-8')
            assert (result.returncode == 0) == ok, result.stdout + result.stderr
            return result
        catalog_path = ROOT/'examples/translation-catalog.json'
        cli('export_workbook.py', catalog_path, '--out', temp/'cli')
        cli('import_workbook.py', catalog_path, temp/'cli/volume-001.xlsx', '--out', temp/'candidate.json')
        assert read_json(temp/'candidate.json') == load_state(adapter)
        candidate_bytes = (temp/'candidate.json').read_bytes()
        cli('import_workbook.py', catalog_path, temp/'cli/volume-001.xlsx', '--out', temp/'candidate.json', ok=False)
        assert (temp/'candidate.json').read_bytes() == candidate_bytes
        # Dedup and splitting: four positions become three units, in two books.
        manifest = export_workbooks(adapter, temp/'split', rows_per_book=2)
        assert [book['units'] for book in manifest['books']] == [2, 1]
        assert sum(book['positions'] for book in manifest['books']) == 4
        conflict = load_state(adapter)
        conflict['languages']['zh-TW'][first] = '歡迎，{player}。'
        conflict['languages']['zh-TW'][second] = '請進，{player}。'
        assert export_workbooks(adapter, temp/'conflict', conflict)['books'][0]['units'] == 4
        export_workbooks(adapter, temp/'中文資料夾')
        workbook = temp/'中文資料夾/volume-001.xlsx'
        map_path = workbook.with_suffix('.map.json')
        pristine = workbook.read_bytes()
        map_bytes = map_path.read_bytes()
        def edit(callback):
            workbook.write_bytes(pristine)
            book = load_workbook(workbook)
            try:
                callback(book)
                book.save(workbook)
            finally:
                book.close()
        book = load_workbook(workbook)
        assert tuple(book.sheetnames) == SHEETS
        assert book[SHEETS[0]]['B4'].value.startswith('=') and book[SHEETS[0]]['B4'].data_type == 's'
        assert book[SHEETS[0]]['B3'].value.count('\n') == 1
        book.close()
        # An empty workbook batch preserves current and unrelated language state.
        assert import_candidate(adapter, workbook, 'zh-TW', state) == state
        cleared = import_candidate(adapter, workbook, 'zh-TW', state, clear_blank=True)
        assert third not in cleared['languages']['zh-TW'] and cleared['languages']['yue-HK'] == state['languages']['yue-HK']
        def translated(book):
            book[SHEETS[0]]['C2'] = '歡迎，{player}。'
            book[SHEETS[2]]['C3'] = '請進，{player}。'
        edit(translated)
        candidate = import_candidate(adapter, workbook, 'zh-TW', state)
        assert candidate['languages']['zh-TW'][first] == '歡迎，{player}。'
        assert candidate['languages']['zh-TW'][second] == '請進，{player}。'
        assert candidate['languages']['zh-TW'][third] == state['languages']['zh-TW'][third]
        assert candidate['languages']['yue-HK'] == state['languages']['yue-HK']
        assert import_candidate(adapter, workbook, 'zh-TW', candidate) == candidate
        assert import_candidate(adapter, temp/'split/volume-002.xlsx', 'zh-TW', candidate,
                                clear_blank=True)['languages']['zh-TW'] == candidate['languages']['zh-TW']
        # All failures leave the old state unchanged and CLI never emits a partial candidate.
        state_before = copy.deepcopy(state)
        wrong_state = copy.deepcopy(state)
        wrong_state['catalog_sha256'] = '0' * 64
        reject(lambda: import_candidate(adapter, workbook, 'zh-TW', wrong_state), 'different original catalog')
        failures = [('C2', '=1+1', 'formulas/numbers'), ('C2', 12, 'formulas/numbers'),
                    ('B2', 'Changed original', 'Source text'), ('C2', 'Missing name', 'Dynamic tokens'),
                    ('C3', 'Missing line break', 'newline'), ('C2', 'Tab\t{player}', 'control')]
        for cell, value, expected in failures:
            edit(lambda book, c=cell, v=value: setattr(book[SHEETS[0]][c], 'value', v))
            reject(lambda: import_candidate(adapter, workbook, 'zh-TW', state), expected)
            cli('import_workbook.py', catalog_path, workbook, '--out', temp/'failed.json', ok=False)
            assert not (temp/'failed.json').exists()
        def two_errors(book):
            book[SHEETS[0]]['C2'] = 'Name missing'
            book[SHEETS[0]]['C3'] = 'Newline missing'
        edit(two_errors)
        try:
            import_candidate(adapter, workbook, 'zh-TW', state)
            raise AssertionError('Expected two row errors')
        except ValueError as error:
            assert '文字對照!2' in str(error) and '文字對照!3' in str(error)
        edit(lambda book: book[SHEETS[0]].append([cell.value for cell in book[SHEETS[0]][2]]))
        reject(lambda: import_candidate(adapter, workbook, 'zh-TW', state), 'Duplicate row')
        edit(lambda book: book[SHEETS[2]].append([cell.value for cell in book[SHEETS[2]][2]]))
        reject(lambda: import_candidate(adapter, workbook, 'zh-TW', state), 'Duplicate row')
        workbook.write_bytes(pristine)
        map_path.rename(temp/'held.map.json')
        reject(lambda: import_candidate(adapter, workbook, 'zh-TW', state), 'Missing companion')
        (temp/'held.map.json').rename(map_path)
        changed = copy.deepcopy(original)
        changed['occurrences'][0]['source'] = 'Different original'
        reject(lambda: import_candidate(SyntheticAdapter(changed), workbook, 'zh-TW'), 'catalog changed')
        mapping = json.loads(map_bytes)
        mapping['units'][0]['occurrences'][0]['id'] = '../escape:T00000'
        mapping['units'][0]['id'] = 'U-' + digest({k: mapping['units'][0][k] for k in ('source', 'mode', 'occurrences')})
        map_path.write_text(json.dumps(mapping), encoding='utf-8')
        reject(lambda: import_candidate(adapter, workbook, 'zh-TW'), 'original position')
        map_path.write_bytes(map_bytes)
        reject(lambda: export_workbooks(adapter, temp/'中文資料夾'))
        assert workbook.read_bytes() == pristine
        # Literal formula-like translated text and whitespace survive untouched.
        def literal_translation(book):
            book[SHEETS[0]]['C4'] = '=純文字'
            book[SHEETS[0]]['C4'].data_type = 's'
            book[SHEETS[0]]['C2'] = '  歡迎，{player}。  '
        edit(literal_translation)
        result = import_candidate(adapter, workbook, 'zh-TW')
        assert result['languages']['zh-TW'][fourth] == '=純文字'
        assert result['languages']['zh-TW'][first] == '  歡迎，{player}。  '
        # ZIP bounds and external relationship/DTD rejection happen before workbook parsing.
        for payload, expected in [
            (b'<!DOCTYPE x [<!ENTITY y "z">]><x/>', 'DTD/entities'),
            (b'<Relationships><Relationship TargetMode="External" Target="https://example.invalid"/></Relationships>', 'External relationship')]:
            workbook.write_bytes(pristine)
            with zipfile.ZipFile(workbook, 'a') as archive:
                archive.writestr('test.rels', payload)
            reject(lambda: import_candidate(adapter, workbook, 'zh-TW'), expected)
        workbook.write_bytes(pristine)
        with zipfile.ZipFile(workbook, 'a', zipfile.ZIP_DEFLATED) as archive:
            archive.writestr('oversize.xml', b' ' * (32 * 1024 * 1024))
        reject(lambda: import_candidate(adapter, workbook, 'zh-TW'), 'Expanded workbook')
        assert adapter.catalog == original and state == state_before
    print(json.dumps({'synthetic_xlsx_roundtrip': True, 'split_and_dedup': True,
                      'existing_conflicts_preserved': True, 'blank_clear_locale_override_repeat': True,
                      'source_map_type_token_errors_rejected': True, 'multiple_errors_have_rows': True,
                      'xml_and_zip_guards': True, 'no_partial_candidate_or_overwrite': True,
                      'game_compilation_or_glyph_tests': False}, indent=2))


if __name__ == '__main__':
    main()
