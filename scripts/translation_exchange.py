"""Bounded XLSX exchange for normalized synthetic text, producing candidate state only."""
import copy
import io
import json
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from engine_adapter import SyntheticAdapter, digest, literal, require

LANGUAGES = ('zh-TW', 'yue-HK', 'zh-CN')
HEADERS = ('ID', 'source', *LANGUAGES)
SHEETS = ('文字對照', '位置對照', '位置覆寫')
MAX_BYTES = 32 * 1024 * 1024


def read_json(path):
    require(path.stat().st_size <= MAX_BYTES, 'JSON exceeds 32 MiB')
    def unique(pairs):
        result = {}
        for key, value in pairs:
            require(key not in result, f'Duplicate JSON key: {key}')
            result[key] = value
        return result
    return json.loads(path.read_text(encoding='utf-8'), object_pairs_hook=unique)


def write_new_json(path, value):
    # Serialize fully before creating output, and never truncate an existing file.
    payload = json.dumps(value, ensure_ascii=False, indent=2) + '\n'
    with path.open('x', encoding='utf-8', newline='\n') as handle:
        handle.write(payload)


def load_state(adapter, state=None):
    if state is None:
        return {'schema': 1, 'catalog_sha256': adapter.fingerprint,
                'languages': {lang: {} for lang in LANGUAGES}}
    require(state.get('schema') == 1 and set(state.get('languages', {})) == set(LANGUAGES),
            'Unknown state schema or languages')
    require(state.get('catalog_sha256') == adapter.fingerprint, 'State belongs to a different original catalog')
    for lang, values in state['languages'].items():
        require(isinstance(values, dict), f'Invalid state language: {lang}')
        for oid, value in values.items():
            require(oid in adapter.occurrences, f'State has unknown position: {oid}')
            adapter.split_translation(adapter.occurrences[oid], value)
    return copy.deepcopy(state)


def make_units(adapter, state):
    groups = {}
    for item in adapter.catalog['occurrences']:
        translations = {lang: state['languages'][lang].get(item['id'], '') for lang in LANGUAGES}
        # Different existing translations must never collapse into one row.
        key = digest({'source': item['source'], 'mode': item['mode'], 'translations': translations})
        group = groups.setdefault(key, {'source': item['source'], 'mode': item['mode'],
                                        'occurrences': [], 'translations': translations})
        group['occurrences'].append(item)
    units = []
    for group in groups.values():
        group['id'] = 'U-' + digest({k: group[k] for k in ('source', 'mode', 'occurrences')})
        units.append(group)
    return units


def append_text(sheet, values):
    row = sheet.max_row + 1 if sheet.cell(1, 1).value is not None else 1
    for column, value in enumerate(values, 1):
        cell = sheet.cell(row, column, value)
        cell.data_type = 's'  # A source beginning with '=' stays text, never a formula.
        cell.number_format = '@'
        cell.alignment = Alignment(vertical='top', wrap_text=True)


def export_workbooks(adapter, destination, state=None, rows_per_book=10000):
    require(1 <= rows_per_book <= 10000, 'rows_per_book must be 1..10000')
    state = load_state(adapter, state)
    units = make_units(adapter, state)
    destination.mkdir(parents=True, exist_ok=False)
    manifest = {'schema': 1, 'catalog_sha256': adapter.fingerprint, 'books': []}
    for start in range(0, len(units), rows_per_book):
        selected = units[start:start + rows_per_book]
        name = f'volume-{len(manifest["books"]) + 1:03}.xlsx'
        book = Workbook()
        book.remove(book.active)
        for title in SHEETS:
            sheet = book.create_sheet(title)
            append_text(sheet, HEADERS if title != SHEETS[1] else ('unit_id', 'occurrence_id', 'source'))
        for unit in selected:
            append_text(book[SHEETS[0]], [unit['id'], unit['source'],
                                          *[unit['translations'][lang] for lang in LANGUAGES]])
            for item in unit['occurrences']:
                append_text(book[SHEETS[1]], [unit['id'], item['id'], item['source']])
                # Overrides start blank so shared edits can take effect.
                append_text(book[SHEETS[2]], [item['id'], item['source'], '', '', ''])
        for sheet in book:
            sheet.freeze_panes = 'C2'
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='244863')
            for column in ('A', 'B', 'C', 'D', 'E'):
                sheet.column_dimensions[column].width = 48 if column != 'A' else 30
        workbook_path = destination/name
        book.save(workbook_path)
        book.close()
        mapping = {'schema': 1, 'adapter': adapter.catalog['adapter'],
                   'catalog_sha256': adapter.fingerprint,
                   'units': [{k: v for k, v in unit.items() if k != 'translations'} for unit in selected]}
        write_new_json(workbook_path.with_suffix('.map.json'), mapping)
        manifest['books'].append({'xlsx': name, 'units': len(selected),
                                  'positions': sum(len(unit['occurrences']) for unit in selected)})
    write_new_json(destination/'manifest.json', manifest)
    return manifest


def checked_workbook(path):
    require(path.suffix.lower() == '.xlsx', 'Select the .xlsx, not its .map.json')
    require(path.stat().st_size <= MAX_BYTES, 'Workbook exceeds 32 MiB')
    # Inspect and parse the same byte snapshot; no race between ZIP check and load.
    data = path.read_bytes()
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        infos = archive.infolist()
        names = [info.filename for info in infos]
        require(len(infos) <= 2000 and len(set(names)) == len(names), 'Excessive/duplicate ZIP entries')
        require(sum(info.file_size for info in infos) <= MAX_BYTES, 'Expanded workbook exceeds 32 MiB')
        for info in infos:
            name = info.filename
            require('..' not in name.split('/') and '\\' not in name and not name.startswith('/'),
                    'Unsafe ZIP member name')
            require(name.endswith(('.xml', '.rels')), 'Only XML workbook parts are supported')
            raw = archive.read(info)
            # This limited reader accepts UTF-8 XML only; reject DTDs before any XML parser.
            xml = raw.decode('utf-8-sig')
            require('<!DOCTYPE' not in xml.upper() and '<!ENTITY' not in xml.upper(), 'DTD/entities rejected')
            root = ET.fromstring(xml)
            for node in root.iter():
                require(node.attrib.get('TargetMode', '').lower() != 'external', 'External relationship rejected')
    book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    try:
        require(tuple(book.sheetnames) == SHEETS, 'Expected the three exported worksheets in order')
        for sheet in book:
            # Do not trust Excel's cached dimension: stream every physical row instead.
            sheet.reset_dimensions()
        return book
    except Exception:
        book.close()
        raise


def validate_mapping(adapter, mapping):
    require(mapping.get('schema') == 1 and mapping.get('adapter') == adapter.catalog['adapter'],
            'Unknown map schema/adapter')
    require(mapping.get('catalog_sha256') == adapter.fingerprint, 'Original catalog changed')
    require(isinstance(mapping.get('units'), list) and 0 < len(mapping['units']) <= 10000, 'Invalid map size')
    units, positions = {}, {}
    for unit in mapping['units']:
        require(set(unit) == {'id', 'source', 'mode', 'occurrences'}, 'Invalid mapped unit fields')
        uid = 'U-' + digest({k: unit[k] for k in ('source', 'mode', 'occurrences')})
        require(uid == unit['id'] and uid not in units, 'Changed/duplicate unit ID')
        require(isinstance(unit['occurrences'], list) and unit['occurrences'], 'Missing mapped positions')
        for item in unit['occurrences']:
            oid = item.get('id')
            require(oid in adapter.occurrences and item == adapter.occurrences[oid], 'Changed original position')
            require(item['source'] == unit['source'] and item['mode'] == unit['mode'], 'Changed unit structure')
            require(oid not in positions, 'Mapped position appears more than once')
            positions[oid] = item
        units[uid] = unit
    return units, positions


def read_rows(sheet, expected_headers, errors):
    rows = sheet.iter_rows()
    first = next(rows, ())
    headers = [cell.value for cell in first]
    require(len(headers) == len(expected_headers) and set(headers) == set(expected_headers),
            f'{sheet.title}: incorrect or duplicate headers')
    require(all(cell.data_type in ('s', 'inlineStr') for cell in first), 'Headers must be text')
    for number, cells in enumerate(rows, 2):
        require(number <= 100001, f'{sheet.title}: row limit exceeded')
        if all(cell.value is None for cell in cells):
            continue
        if len(cells) > len(headers) and any(c.value is not None for c in cells[len(headers):]):
            errors.append(f'{sheet.title}!{number}: unexpected extra cells')
            continue
        values = {}
        valid = True
        for index, header in enumerate(headers):
            cell = cells[index] if index < len(cells) else None
            value = cell.value if cell else None
            if value is not None and (not isinstance(value, str) or cell.data_type not in ('s', 'inlineStr')):
                errors.append(f'{sheet.title}!{number} ({header}): text required; formulas/numbers rejected')
                valid = False
            values[header] = '' if value is None else value
        if valid:
            yield number, values


def import_candidate(adapter, workbook_path, language, state=None, clear_blank=False):
    require(language in LANGUAGES, 'Unknown language')
    require(workbook_path.suffix.lower() == '.xlsx', 'Select an .xlsx workbook')
    mapping_path = workbook_path.with_suffix('.map.json')
    require(mapping_path.is_file(), 'Missing companion .map.json')
    units, positions = validate_mapping(adapter, read_json(mapping_path))
    candidate = load_state(adapter, state)
    updates, overrides, errors = {}, {}, []
    book = checked_workbook(workbook_path)
    try:
        for title, table in ((SHEETS[0], units), (SHEETS[2], positions)):
            seen = set()
            for number, row in read_rows(book[title], HEADERS, errors):
                where = f'{title}!{number}'
                try:
                    identifier = row['ID']
                    require(identifier in table, 'Unknown ID')
                    require(identifier not in seen, 'Duplicate row ID')
                    seen.add(identifier)
                    item = table[identifier]
                    require(row['source'] == item['source'], 'Source text changed')
                    value = row[language]
                    targets = item['occurrences'] if title == SHEETS[0] else [item]
                    # Blank overrides always mean no override; clearing is a unit-row action.
                    if value == '' and (not clear_blank or title == SHEETS[2]):
                        continue
                    for occurrence in targets:
                        if value != '':
                            parts = adapter.split_translation(occurrence, value)
                            require(len(parts) == 1 and parts[0] == value, 'Synthetic split identity failed')
                        destination = updates if title == SHEETS[0] else overrides
                        require(occurrence['id'] not in destination, 'Duplicate expanded position')
                        destination[occurrence['id']] = value
                except ValueError as error:
                    errors.append(f'{where}: {error}')
        # Reference-only rows are checked for text types, but never used as write targets.
        list(read_rows(book[SHEETS[1]], ('unit_id', 'occurrence_id', 'source'), errors))
    finally:
        book.close()
    require(not errors, '\n'.join(errors[:50]) + (f'\nTotal errors: {len(errors)}' if len(errors) > 50 else ''))
    updates.update(overrides)
    for oid, value in updates.items():
        if value == '':
            candidate['languages'][language].pop(oid, None)
        else:
            candidate['languages'][language][oid] = value
    return candidate
